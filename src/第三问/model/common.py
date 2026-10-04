"""Portable constants/IO; the numerical model constants are unchanged."""
from pathlib import Path
import os,json,hashlib,time,csv
os.environ.setdefault('OMP_NUM_THREADS','1')
os.environ.setdefault('OPENBLAS_NUM_THREADS','1')
import numpy as np
PACKAGE=Path(__file__).resolve().parents[1]
LOW,HIGH,ETA,B=1200.,10800.,.9,5000./6.
DATES=np.arange(np.datetime64('2025-01-01'),np.datetime64('2026-01-01'))
def digest(x):return hashlib.sha256(np.asarray(x).tobytes()).hexdigest()
def read_npz(path):
    with np.load(path,allow_pickle=False) as z:return {k:z[k].copy() for k in z.files}
def filehash(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def savejson(path,value):
    def encode(x):
        if isinstance(x,np.ndarray):return x.tolist()
        if isinstance(x,np.generic):return x.item()
        if isinstance(x,Path):return str(x)
        raise TypeError(type(x).__name__)
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_suffix(path.suffix+'.tmp')
    tmp.write_text(json.dumps(value,ensure_ascii=False,indent=2,default=encode),encoding='utf-8');tmp.replace(path)
def inputs():
    from openpyxl import load_workbook
    attachment=PACKAGE.parent/'附件'
    wb=load_workbook(attachment/'附件1.xlsx',read_only=True,data_only=True)
    p=np.asarray([r[1] for r in list(wb.worksheets[0].values)[1:]],float);wb.close()
    wb=load_workbook(attachment/'附件2.xlsx',read_only=True,data_only=True)
    data=[]
    for sheet in wb.worksheets[:2]:
        rows=list(sheet.values)
        assert [str(r[0].date()) for r in rows[1:]]==list(map(str,DATES))
        data.append(np.asarray([r[1:145] for r in rows[1:]],float)/6.)
    wb.close();L,V=data
    assert L.shape==V.shape==(365,144) and p.shape==(144,)
    assert np.isfinite(L).all() and np.isfinite(V).all() and np.isfinite(p).all() and p.min()>0
    return dict(L=L,V=V,p=p)
