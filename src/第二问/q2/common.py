"""Physical constants, package-local data and portable file helpers.

All scheduling quantities are ten-minute AC-side energy in kWh.
SOC variables are battery energy in kWh, not fractions.
"""
from pathlib import Path
import os, json, hashlib, time, csv
os.environ.setdefault('OMP_NUM_THREADS', '1')
os.environ.setdefault('OPENBLAS_NUM_THREADS', '1')
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
LOW, HIGH, ETA, B = 1200.0, 10800.0, 0.9, 5000.0 / 6
INITIAL_SOC = 6000.0
DATES = np.arange(np.datetime64('2025-01-01'), np.datetime64('2026-01-01'))
with np.load(DATA / 'calendar.npz') as _calendar:
    CAL = dict(_calendar)

def savejson(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    def convert(value):
        if isinstance(value, np.ndarray): return value.tolist()
        if isinstance(value, np.generic): return value.item()
        return str(value)
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=convert), encoding='utf-8')
    tmp.replace(path)

def digest(x):
    return hashlib.sha256(np.asarray(x).tobytes()).hexdigest()

def save_npz(path, **arrays):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix('.tmp')
    with tmp.open('wb') as f: np.savez_compressed(f, **arrays)
    tmp.replace(path)

def csvout(path, rows):
    rows = list(rows)
    if not rows: return
    keys = list(dict.fromkeys(k for row in rows for k in row))
    with Path(path).open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=keys)
        writer.writeheader(); writer.writerows(rows)

def inputs():
    """Read original attachments. Samples are right endpoints of ten-minute slots."""
    from openpyxl import load_workbook
    wb = load_workbook(DATA / 'attachment1.xlsx', read_only=True, data_only=True)
    rows = list(wb.worksheets[0].values); wb.close()
    p = np.array([r[1] for r in rows[1:]], float)
    wb = load_workbook(DATA / 'attachment2.xlsx', read_only=True, data_only=True)
    load = list(wb.worksheets[0].values); pv = list(wb.worksheets[1].values); wb.close()
    L = np.array([r[1:] for r in load[1:]], float) / 6
    V = np.array([r[1:] for r in pv[1:]], float) / 6
    assert L.shape == V.shape == (365, 144) and p.shape == (144,)
    assert np.isfinite(L).all() and np.isfinite(V).all() and np.all(p > 0)
    assert [str(r[0].date()) for r in load[1:]] == [str(d) for d in DATES]
    assert [r[0] for r in load] == [r[0] for r in pv]
    return L, V, p

def verify_data():
    ledger = json.loads((DATA / 'sha256.json').read_text(encoding='utf-8'))
    for name, expected in ledger.items():
        actual = hashlib.sha256((ROOT / name).read_bytes()).hexdigest()
        if actual != expected: raise ValueError('Input checksum mismatch: ' + name)
    return ledger
