"""Q4-2: daily causal price LightGBM + unchanged net-load Q80 and policy B."""
import os,sys
from pathlib import Path
WORK=Path(__file__).resolve().parents[1]
SOURCE=WORK.parent/'第二问最终版_Q80'
os.environ['OMP_NUM_THREADS']='1';os.environ['OPENBLAS_NUM_THREADS']='1'
sys.path[:0]=[str(WORK/'.deps'),str(SOURCE)]
from q2.common import *
from q2.final_pipeline import make
from q2.features import naive
from q2.settlement import dispatch
from q2.risk_terminal import terminal_optimize
from q4_value import continuation
from lightgbm import LGBMRegressor
from openpyxl import load_workbook
import argparse

PARAMS=dict(objective='regression',n_estimators=120,learning_rate=.04,num_leaves=15,
    min_child_samples=144,reg_lambda=20.,reg_alpha=.1,verbosity=-1,n_jobs=1,random_state=20260913)

def price_features(history,lead):
    """No actual value at or after issuance is accepted by this interface."""
    d=len(history);target=d+lead;t=np.arange(144);cols=[t/144,np.full(144,lead)]
    for k in range(1,5):cols.extend([np.sin(2*np.pi*k*t/144),np.cos(2*np.pi*k*t/144)])
    cols.extend([np.full(144,(target+2)%7),np.full(144,np.sin(2*np.pi*target/365)),np.full(144,np.cos(2*np.pi*target/365))])
    for lag in [1,2,3,7,14,28]:
        idx=target-lag
        cols.append(history[idx] if 0<=idx<d else history[-1])
    for win in [3,7,14,28]:
        h=history[-win:];cols.extend([h.mean(0),h.std(0)])
    cols.extend([np.full(144,history[-1].mean()),history[-1]-history[-min(7,d)]])
    return np.column_stack(cols)

def price_issue(history,prior,cached=None):
    d=len(history)
    if d<21:
        p=history[-3:].mean(0) if d else prior.copy()
        return np.array([p,p]),dict(method='cold_start_prior' if not d else 'past3_mean',train_last=d-1)
    # All labels, including lead=1 labels, end strictly before issuance day d.
    pairs=[(i,h) for i in range(max(7,d-180),d) for h in [0,1] if i+h<d]
    X=np.concatenate([cached[i,h] if cached is not None else price_features(history[:i],h) for i,h in pairs])
    y=np.concatenate([history[i+h] for i,h in pairs])
    reg=LGBMRegressor(**PARAMS).fit(X,y)
    pred=np.array([reg.predict(price_features(history,h)) for h in [0,1]])
    return np.maximum(pred,1e-6),dict(method='direct_LightGBM',train_first=pairs[0][0],train_last=d-1,
        latest_label_day=max(i+h for i,h in pairs),training_pairs=len(pairs),history_hash=digest(history))

def load_inputs():
    # Q4 has no attachment1 input. Read measured load/PV directly from attachment2.
    wb=load_workbook(WORK.parent/'附件/附件2.xlsx',read_only=True,data_only=True)
    arrays=[]
    for sheet in wb.worksheets[:2]:
        rows=list(sheet.values)
        assert [str(r[0].date()) for r in rows[1:]]==list(map(str,DATES))
        arrays.append(np.array([r[1:] for r in rows[1:]],float)/6)
    wb.close()
    L,V=arrays
    assert L.shape==V.shape==(365,144) and np.isfinite(L).all() and np.isfinite(V).all()
    wb=load_workbook(WORK.parent/'附件/附件4.xlsx',read_only=True,data_only=True)
    rows=list(wb.active.values);wb.close()
    prices=np.array([r[1:] for r in rows[1:]],float)
    assert prices.shape==(365,144) and np.isfinite(prices).all() and prices.min()>0
    assert [str(r[0].date()) for r in rows[1:]]==list(map(str,DATES))
    # User-defined initial information: Jan1 attachment4 curve is known at startup.
    # From Jan2 onwards, forecasting receives completed days only.
    prior=prices[0].copy()
    ledger=json.loads((DATA/'sha256.json').read_text(encoding='utf-8'))
    for name in ['data/issued_forecasts.npz','data/calendar.npz']:
        assert hashlib.sha256((SOURCE/name).read_bytes()).hexdigest()==ledger[name]
    with np.load(DATA/'issued_forecasts.npz') as z:issued=dict(z)
    return L,V,prior,prices,issued

def net_issue(Lpast,Vpast,issued,d,prior):
    if d<21:
        return (naive(Lpast,True)-naive(Vpast,False))[None],None
    # Slice both actuals and issuance archives at this decision boundary.
    arc={k:v[:d+1] for k,v in issued.items()}
    n0,n1,_=make(Lpast,Vpast,prior,d,'q80',arc)
    return n0,n1

def procure(n0,n1,pred,soc,d):
    if d<21 or d==364:
        q,st,_=terminal_optimize(n0,pred[0],soc,[],limit=20,gap=.0001)
        de=None
    else:
        q,st,de=continuation(n0,n1,pred[0],soc,pnext=pred[1])
        if st['gap']>.0001+1e-8:
            q,st,de=continuation(n0,n1,pred[0],soc,pnext=pred[1],limit=60,max_iter=15)
    assert st['gap']<=.0001+1e-8, (d,st['gap'])
    # Independently replay predicted objective, including tomorrow's own prices.
    z=dispatch(q,n0[0],soc);objective=float(pred[0]@(q+5*z['e']))
    if de is not None:
        qt=de['provisional_q'][0];zt=dispatch(qt,n1[0],z['s'][-1])
        objective+=float(pred[1]@(qt+5*zt['e']))
    assert abs(objective-st['incumbent'])<1e-4,(objective,st['incumbent'])
    return q,st

def leakage_tests(L,V,prior,P,issued):
    records=[]
    for d in [1,21,31,180,364]:
        p0,_=price_issue(P[:d],prior)
        bad=P.copy();bad[d:]=12345
        p1,_=price_issue(bad[:d],prior)
        assert np.array_equal(p0,p1)
        a,b=net_issue(L[:d],V[:d],issued,d,prior)
        ll=L.copy();vv=V.copy();ll[d:]=1e9;vv[d:]=1e8
        arc={k:v.copy() for k,v in issued.items()}
        for v in arc.values():v[d+1:]=1e9
        aa,bb=net_issue(ll[:d],vv[:d],arc,d,prior)
        assert np.array_equal(a,aa) and np.array_equal(b,bb)
        # Deterministic, certified optimisation on unchanged inputs must replay.
        q0,_=procure(a,b,p0,6000,d);q1,_=procure(aa,bb,p1,6000,d)
        assert np.allclose(q0,q1,atol=1e-7,rtol=0)
        records.append(dict(day=str(DATES[d]),price_forecast_unchanged=True,net_forecast_unchanged=True,
                            procurement_max_difference=float(np.max(abs(q0-q1)))))
    return records

def code_manifest():
    paths=sorted(Path(__file__).parent.glob('*.py'))+sorted((SOURCE/'q2').glob('*.py'))
    return {p.relative_to(WORK.parent).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--smoke',action='store_true');args=ap.parse_args()
    out=WORK/('试跑' if args.smoke else '结果4-2');out.mkdir(parents=True,exist_ok=True)
    L,V,prior,P,issued=load_inputs();days=33 if args.smoke else 365
    cfg=dict(version=3,days=days,source_hashes=code_manifest(),price_model=PARAMS,net_model='frozen Q2 ordinary LightGBM + residual Q80',
        information='Jan1 attachment4 price curve is an explicitly known startup input; Jan2 onwards uses completed days only',
        price_prior='attachment4 Jan1 curve on Jan1; last<=3 completed days from Jan2 through Jan21',
        settlement='delivery-interval actual price times (planned kWh + 5 * emergency kWh)',
        startup='recomputed from Jan1 SOC6000 under predicted prices',
        price_hash=digest(P),load_hash=digest(L),pv_hash=digest(V),issued_hash=digest(issued['today']))
    if (out/'config.json').exists():assert json.loads((out/'config.json').read_text(encoding='utf-8'))==cfg
    savejson(out/'config.json',cfg)
    cache={(i,h):price_features(P[:i],h) for i in range(7,days) for h in [0,1]}
    a={k:np.zeros((days,145 if k=='s' else 144)) for k in ['q','c','d','e','w','s']}
    forecasts=np.zeros((days,2,144));logs=[];begin=0;cp=out/'checkpoint.npz'
    if cp.exists():
        with np.load(cp) as z:
            for k in a:a[k]=z[k].copy()
            forecasts=z['price_forecasts'].copy();logs=json.loads(str(z['logs']));begin=len(logs)
    start=time.time()
    for d in range(begin,days):
        pred,meta=price_issue(P[:d],prior,cache);forecasts[d]=pred
        n0,n1=net_issue(L[:d],V[:d],issued,d,prior)
        soc=6000. if d==0 else float(a['s'][d-1,-1])
        q,st=procure(n0,n1,pred,soc,d)
        a['q'][d]=q
        # Only after q is locked may this function see current-day actual load/PV.
        for k,v in dispatch(q,L[d]-V[d],soc).items():a[k][d]=v
        logs.append(dict(day=d,date=str(DATES[d]),price=meta,forecast_hash=digest(pred),q_hash=digest(q),solver=st))
        save_npz(cp,**a,price_forecasts=forecasts,logs=np.array(json.dumps(logs)))
        if d%10==0 or d==days-1:print(str(DATES[d]),f'{d+1}/{days}',round(time.time()-start,1),'seconds',flush=True)
    balance=float(np.max(abs(a['q']+a['d']+a['e']-(L[:days]-V[:days])-a['c']-a['w'])))
    socerr=float(np.max(abs(np.diff(a['s'],axis=1)-.9*a['c']+a['d']/.9)))
    assert balance<1e-7 and socerr<1e-7 and np.min(a['s'])>=1200-1e-7 and np.max(a['s'])<=10800+1e-7
    assert np.allclose(a['s'][1:,0],a['s'][:-1,-1],atol=1e-8)
    assert max(a['c'].max(),a['d'].max())<=5000/6+1e-7 and np.max(a['c']*a['d'])<1e-7
    for d in range(days):
        for k,v in dispatch(a['q'][d],L[d]-V[d],a['s'][d,0]).items():assert np.allclose(v,a[k][d],atol=1e-7)
    normal=(P[:days]*a['q']).sum(1);emergency=5*(P[:days]*a['e']).sum(1)
    summary=dict(formal_days=days-31,normal_cost=float(normal[31:].sum()),emergency_cost=float(emergency[31:].sum()),
        total_cost=float((normal+emergency)[31:].sum()),emergency_kwh=float(a['e'][31:].sum()),
        january_cost=float((normal+emergency)[:31].sum()),initial_formal_soc=float(a['s'][31,0]),
        final_soc=float(a['s'][-1,-1]),price_MAE=float(abs(forecasts[31:,0]-P[31:days]).mean()),
        gap_max=max(x['solver']['gap'] for x in logs[31:]))
    savejson(out/'summary.json',summary)
    csvout(out/'daily.csv',[dict(date=str(DATES[d]),normal_cost=normal[d],emergency_cost=emergency[d],
        total_cost=normal[d]+emergency[d],emergency_kwh=a['e'][d].sum(),initial_soc=a['s'][d,0],final_soc=a['s'][d,-1]) for d in range(31,days)])
    tests=leakage_tests(L,V,prior,P,issued)
    assert np.array_equal(forecasts[0,0],P[0])
    savejson(out/'audit.json',dict(status='PASS',energy_balance_error=balance,soc_transition_error=socerr,future_poison_tests=tests,
        no_attachment1_used=True,jan1_price_curve_explicitly_known=True,jan1_prior_matches_attachment4=True,
        no_attachment3_used=True,all_day_price_training_cutoffs_valid=all(x['price']['train_last']==x['day']-1 for x in logs)))
    save_npz(out/'q4_2.npz',**a,price_forecasts=forecasts)
    # Plain data for the template exporter; all formal intervals retained.
    if not args.smoke:savejson(out/'export.json',dict(dates=list(map(str,DATES[31:])),**{k:v[31:] for k,v in a.items()},daily_cost=(normal+emergency)[31:]))
    print(json.dumps(summary,ensure_ascii=False),flush=True)

if __name__=='__main__':main()
